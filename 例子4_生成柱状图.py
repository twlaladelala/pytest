from openpyxl import Workbook
from openpyxl.chart import BarChart, Series, Reference

wb = Workbook(write_only=True)
ws = wb.create_sheet()

rows = [
    ('类别', '销售A组', '销售B组'),
    ('手机', 10, 30),
    ('平板', 40, 60),
    ('笔记本', 50, 70),
    ('外围设备', 20, 10),
    ('铅笔', 10, 40),
    ('相机', 50, 30),
]

for row in rows:
    ws.append(row)

chart1 = BarChart()  # 柱状图
# 柱状图类型 ，样式，标题
chart1.type = 'col'
chart1.style = 10
chart1.title = '销售统计图'
chart1.y_axis.title = '销量'
chart1.x_axis.title = '商品类别'

data = Reference(ws, min_col=2, min_row=1, max_row=7, max_col=3)
cats = Reference(ws, min_col=1, min_row=2, max_row=7)  # 分类数据
print(data)
print(cats)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
ws.add_chart(chart1, 'A10')


wb.save('table/销售统计图demo.xlsx')
